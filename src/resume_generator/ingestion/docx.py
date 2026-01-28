"""DOCX text extraction using python-docx."""

from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from resume_generator.ingestion.base import BaseExtractor, ExtractionError, ExtractionResult


class DocxExtractor(BaseExtractor):
    """Extracts text content from DOCX files using python-docx."""

    @property
    def supported_extensions(self) -> frozenset[str]:
        return frozenset({".docx"})

    def extract(self, path: Path) -> str:
        """Extract text from a DOCX file.

        Args:
            path: Path to the DOCX file.

        Returns:
            Extracted text with paragraphs separated by newlines.

        Raises:
            ExtractionError: If DOCX cannot be read or parsed.
            FileNotFoundError: If the file does not exist.
        """
        self._validate_path(path)
        try:
            doc = Document(str(path))
            paragraphs: list[str] = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)

            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))

            return "\n\n".join(paragraphs)
        except PackageNotFoundError as e:
            raise ExtractionError(path, f"Invalid or corrupted DOCX file: {e}") from e
        except Exception as e:
            raise ExtractionError(path, str(e)) from e

    def extract_with_metadata(self, path: Path) -> ExtractionResult:
        """Extract text and DOCX metadata.

        Args:
            path: Path to the DOCX file.

        Returns:
            ExtractionResult with text and document properties.
        """
        self._validate_path(path)
        try:
            doc = Document(str(path))
            paragraphs: list[str] = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    paragraphs.append(text)

            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))

            text = "\n\n".join(paragraphs)
            metadata: dict[str, str] = {
                "paragraph_count": str(len(doc.paragraphs)),
                "table_count": str(len(doc.tables)),
            }

            core_props = doc.core_properties
            if core_props.title:
                metadata["title"] = core_props.title
            if core_props.author:
                metadata["author"] = core_props.author
            if core_props.subject:
                metadata["subject"] = core_props.subject

            return ExtractionResult(text=text, source_path=path, metadata=metadata)
        except PackageNotFoundError as e:
            raise ExtractionError(path, f"Invalid or corrupted DOCX file: {e}") from e
        except Exception as e:
            raise ExtractionError(path, str(e)) from e
