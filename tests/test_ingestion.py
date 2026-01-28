"""Unit tests for ingestion module."""

from pathlib import Path

import pytest

from resume_generator.ingestion import (
    DataLoader,
    DocxExtractor,
    ExtractionError,
    ExtractionResult,
    LoadResult,
    PDFExtractor,
    TextExtractor,
)


class TestPDFExtractor:
    """Tests for PDFExtractor."""

    @pytest.fixture
    def extractor(self) -> PDFExtractor:
        return PDFExtractor()

    @pytest.fixture
    def sample_pdf(self, tmp_path: Path) -> Path:
        pdf_file = tmp_path / "test.pdf"
        pdf_file.write_bytes(
            b"%PDF-1.4\n1 0 obj\n<<\n/Type /Catalog\n/Pages 2 0 R\n>>\nendobj\n"
            b"2 0 obj\n<<\n/Type /Pages\n/Kids [3 0 R]\n/Count 1\n>>\nendobj\n"
            b"3 0 obj\n<<\n/Type /Page\n/Parent 2 0 R\n/MediaBox [0 0 612 792]\n"
            b"/Contents 4 0 R\n/Resources <<\n/Font <<\n/F1 <<\n/Type /Font\n"
            b"/Subtype /Type1\n/BaseFont /Helvetica\n>>\n>>\n>>\n>>\nendobj\n"
            b"4 0 obj\n<<\n/Length 44\n>>\nstream\nBT\n/F1 12 Tf\n100 700 Td\n"
            b"(Test PDF) Tj\nET\nendstream\nendobj\nxref\n0 5\n"
            b"0000000000 65535 f\n0000000009 00000 n\n0000000058 00000 n\n"
            b"0000000115 00000 n\n0000000267 00000 n\ntrailer\n<<\n"
            b"/Size 5\n/Root 1 0 R\n>>\nstartxref\n359\n%%EOF\n"
        )
        return pdf_file

    def test_supported_extensions(self, extractor: PDFExtractor) -> None:
        assert ".pdf" in extractor.supported_extensions
        assert len(extractor.supported_extensions) == 1

    def test_can_handle_pdf(self, extractor: PDFExtractor, tmp_path: Path) -> None:
        pdf_path = tmp_path / "test.pdf"
        pdf_path.touch()
        assert extractor.can_handle(pdf_path)

    def test_cannot_handle_non_pdf(self, extractor: PDFExtractor, tmp_path: Path) -> None:
        txt_path = tmp_path / "test.txt"
        txt_path.touch()
        assert not extractor.can_handle(txt_path)

    def test_extract_from_pdf(self, extractor: PDFExtractor, sample_pdf: Path) -> None:
        text = extractor.extract(sample_pdf)
        assert isinstance(text, str)
        assert "Test PDF" in text

    def test_extract_file_not_found(self, extractor: PDFExtractor, tmp_path: Path) -> None:
        nonexistent = tmp_path / "nonexistent.pdf"
        with pytest.raises(FileNotFoundError):
            extractor.extract(nonexistent)

    def test_extract_unsupported_extension(self, extractor: PDFExtractor, tmp_path: Path) -> None:
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("content")
        with pytest.raises(ExtractionError) as exc_info:
            extractor.extract(txt_file)
        assert "Unsupported file type" in str(exc_info.value)

    def test_extract_invalid_pdf(self, extractor: PDFExtractor, tmp_path: Path) -> None:
        invalid_pdf = tmp_path / "invalid.pdf"
        invalid_pdf.write_text("not a pdf")
        with pytest.raises(ExtractionError):
            extractor.extract(invalid_pdf)

    def test_extract_with_metadata(self, extractor: PDFExtractor, sample_pdf: Path) -> None:
        result = extractor.extract_with_metadata(sample_pdf)
        assert isinstance(result, ExtractionResult)
        assert result.source_path == sample_pdf
        assert "Test PDF" in result.text
        assert result.char_count > 0
        assert result.word_count > 0
        assert "page_count" in result.metadata

    def test_extract_directory_raises_error(self, extractor: PDFExtractor, tmp_path: Path) -> None:
        with pytest.raises(ExtractionError) as exc_info:
            extractor.extract(tmp_path)
        assert "not a file" in str(exc_info.value)


class TestTextExtractor:
    """Tests for TextExtractor."""

    @pytest.fixture
    def extractor(self) -> TextExtractor:
        return TextExtractor()

    def test_supported_extensions(self, extractor: TextExtractor) -> None:
        expected = {".txt", ".md", ".markdown", ".text"}
        assert extractor.supported_extensions == expected

    def test_can_handle_text_files(self, extractor: TextExtractor, tmp_path: Path) -> None:
        txt_file = tmp_path / "test.txt"
        txt_file.touch()
        assert extractor.can_handle(txt_file)

        md_file = tmp_path / "test.md"
        md_file.touch()
        assert extractor.can_handle(md_file)

        markdown_file = tmp_path / "test.markdown"
        markdown_file.touch()
        assert extractor.can_handle(markdown_file)

    def test_cannot_handle_non_text(self, extractor: TextExtractor, tmp_path: Path) -> None:
        pdf_path = tmp_path / "test.pdf"
        pdf_path.touch()
        assert not extractor.can_handle(pdf_path)

    def test_extract_from_txt(self, extractor: TextExtractor, tmp_path: Path) -> None:
        txt_file = tmp_path / "test.txt"
        content = "Hello World\nThis is a test file."
        txt_file.write_text(content, encoding="utf-8")

        text = extractor.extract(txt_file)
        assert text.strip() == content.strip()

    def test_extract_from_markdown(self, extractor: TextExtractor, tmp_path: Path) -> None:
        md_file = tmp_path / "test.md"
        content = "# Title\n\nThis is **markdown** content."
        md_file.write_text(content, encoding="utf-8")

        text = extractor.extract(md_file)
        assert text.strip() == content.strip()

    def test_extract_strips_whitespace(self, extractor: TextExtractor, tmp_path: Path) -> None:
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("\n\n  Hello World  \n\n", encoding="utf-8")

        text = extractor.extract(txt_file)
        assert text == "Hello World"

    def test_extract_file_not_found(self, extractor: TextExtractor, tmp_path: Path) -> None:
        nonexistent = tmp_path / "nonexistent.txt"
        with pytest.raises(FileNotFoundError):
            extractor.extract(nonexistent)

    def test_extract_unsupported_extension(self, extractor: TextExtractor, tmp_path: Path) -> None:
        pdf_file = tmp_path / "test.pdf"
        pdf_file.write_text("content")
        with pytest.raises(ExtractionError) as exc_info:
            extractor.extract(pdf_file)
        assert "Unsupported file type" in str(exc_info.value)

    def test_extract_latin1_encoding(self, extractor: TextExtractor, tmp_path: Path) -> None:
        txt_file = tmp_path / "test.txt"
        txt_file.write_bytes("Café".encode("latin-1"))

        text = extractor.extract(txt_file)
        assert "Caf" in text

    def test_extract_with_metadata(self, extractor: TextExtractor, tmp_path: Path) -> None:
        txt_file = tmp_path / "test.txt"
        content = "Test content"
        txt_file.write_text(content, encoding="utf-8")

        result = extractor.extract_with_metadata(txt_file)
        assert isinstance(result, ExtractionResult)
        assert result.text == content
        assert result.source_path == txt_file
        assert result.char_count == len(content)
        assert result.word_count == 2
        assert result.metadata["file_type"] == "txt"
        assert "file_size_bytes" in result.metadata

    def test_from_raw_text(self, extractor: TextExtractor) -> None:
        raw_text = "  Some raw text  "
        result = TextExtractor.from_raw_text(raw_text)

        assert isinstance(result, ExtractionResult)
        assert result.text == "Some raw text"
        assert result.metadata["source_type"] == "raw_text"
        assert result.source_path == Path("raw_input")

    def test_from_raw_text_custom_source_name(self, extractor: TextExtractor) -> None:
        raw_text = "Custom text"
        result = TextExtractor.from_raw_text(raw_text, source_name="custom_source")

        assert result.source_path == Path("custom_source")

    def test_extract_directory_raises_error(self, extractor: TextExtractor, tmp_path: Path) -> None:
        with pytest.raises(ExtractionError) as exc_info:
            extractor.extract(tmp_path)
        assert "not a file" in str(exc_info.value)


class TestDocxExtractor:
    """Tests for DocxExtractor."""

    @pytest.fixture
    def extractor(self) -> DocxExtractor:
        return DocxExtractor()

    @pytest.fixture
    def sample_docx(self, tmp_path: Path) -> Path:
        from docx import Document

        docx_file = tmp_path / "test.docx"
        doc = Document()
        doc.add_paragraph("Test Document Title")
        doc.add_paragraph("This is the first paragraph with some content.")
        doc.add_paragraph("This is the second paragraph.")
        doc.save(docx_file)
        return docx_file

    @pytest.fixture
    def docx_with_table(self, tmp_path: Path) -> Path:
        from docx import Document

        docx_file = tmp_path / "table.docx"
        doc = Document()
        doc.add_paragraph("Document with Table")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header 1"
        table.cell(0, 1).text = "Header 2"
        table.cell(1, 0).text = "Value 1"
        table.cell(1, 1).text = "Value 2"
        doc.save(docx_file)
        return docx_file

    def test_supported_extensions(self, extractor: DocxExtractor) -> None:
        assert ".docx" in extractor.supported_extensions
        assert len(extractor.supported_extensions) == 1

    def test_can_handle_docx(self, extractor: DocxExtractor, tmp_path: Path) -> None:
        docx_path = tmp_path / "test.docx"
        docx_path.touch()
        assert extractor.can_handle(docx_path)

    def test_cannot_handle_non_docx(self, extractor: DocxExtractor, tmp_path: Path) -> None:
        txt_path = tmp_path / "test.txt"
        txt_path.touch()
        assert not extractor.can_handle(txt_path)

        pdf_path = tmp_path / "test.pdf"
        pdf_path.touch()
        assert not extractor.can_handle(pdf_path)

    def test_extract_from_docx(self, extractor: DocxExtractor, sample_docx: Path) -> None:
        text = extractor.extract(sample_docx)
        assert isinstance(text, str)
        assert "Test Document Title" in text
        assert "first paragraph" in text
        assert "second paragraph" in text

    def test_extract_with_table(self, extractor: DocxExtractor, docx_with_table: Path) -> None:
        text = extractor.extract(docx_with_table)
        assert "Document with Table" in text
        assert "Header 1" in text
        assert "Value 1" in text

    def test_extract_file_not_found(self, extractor: DocxExtractor, tmp_path: Path) -> None:
        nonexistent = tmp_path / "nonexistent.docx"
        with pytest.raises(FileNotFoundError):
            extractor.extract(nonexistent)

    def test_extract_unsupported_extension(self, extractor: DocxExtractor, tmp_path: Path) -> None:
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("content")
        with pytest.raises(ExtractionError) as exc_info:
            extractor.extract(txt_file)
        assert "Unsupported file type" in str(exc_info.value)

    def test_extract_invalid_docx(self, extractor: DocxExtractor, tmp_path: Path) -> None:
        invalid_docx = tmp_path / "invalid.docx"
        invalid_docx.write_text("not a docx")
        with pytest.raises(ExtractionError):
            extractor.extract(invalid_docx)

    def test_extract_with_metadata(self, extractor: DocxExtractor, sample_docx: Path) -> None:
        result = extractor.extract_with_metadata(sample_docx)
        assert isinstance(result, ExtractionResult)
        assert result.source_path == sample_docx
        assert "Test Document Title" in result.text
        assert result.char_count > 0
        assert result.word_count > 0
        assert "paragraph_count" in result.metadata
        assert "table_count" in result.metadata

    def test_extract_directory_raises_error(self, extractor: DocxExtractor, tmp_path: Path) -> None:
        with pytest.raises(ExtractionError) as exc_info:
            extractor.extract(tmp_path)
        assert "not a file" in str(exc_info.value)


class TestExtractionResult:
    """Tests for ExtractionResult dataclass."""

    def test_create_result(self, tmp_path: Path) -> None:
        result = ExtractionResult(
            text="Hello World",
            source_path=tmp_path / "test.txt",
        )
        assert result.text == "Hello World"
        assert result.char_count == 11
        assert result.word_count == 2
        assert result.metadata == {}

    def test_result_with_metadata(self, tmp_path: Path) -> None:
        result = ExtractionResult(
            text="Test",
            source_path=tmp_path / "test.txt",
            metadata={"key": "value"},
        )
        assert result.metadata["key"] == "value"

    def test_is_empty_property(self, tmp_path: Path) -> None:
        empty = ExtractionResult(text="", source_path=tmp_path / "empty.txt")
        assert empty.is_empty

        non_empty = ExtractionResult(text="content", source_path=tmp_path / "test.txt")
        assert not non_empty.is_empty

    def test_char_and_word_count(self, tmp_path: Path) -> None:
        text = "This is a test with 123 numbers"
        result = ExtractionResult(text=text, source_path=tmp_path / "test.txt")
        assert result.char_count == len(text)
        assert result.word_count == 7


class TestDataLoader:
    """Tests for DataLoader."""

    @pytest.fixture
    def loader(self) -> DataLoader:
        return DataLoader()

    @pytest.fixture
    def sample_text_file(self, tmp_path: Path) -> Path:
        txt_file = tmp_path / "sample.txt"
        txt_file.write_text("Sample text content", encoding="utf-8")
        return txt_file

    @pytest.fixture
    def sample_md_file(self, tmp_path: Path) -> Path:
        md_file = tmp_path / "sample.md"
        md_file.write_text("# Markdown\n\nContent here", encoding="utf-8")
        return md_file

    @pytest.fixture
    def sample_pdf_file(self, tmp_path: Path) -> Path:
        pdf_file = tmp_path / "sample.pdf"
        pdf_file.write_bytes(
            b"%PDF-1.4\n1 0 obj\n<<\n/Type /Catalog\n/Pages 2 0 R\n>>\nendobj\n"
            b"2 0 obj\n<<\n/Type /Pages\n/Kids [3 0 R]\n/Count 1\n>>\nendobj\n"
            b"3 0 obj\n<<\n/Type /Page\n/Parent 2 0 R\n/MediaBox [0 0 612 792]\n"
            b"/Contents 4 0 R\n/Resources <<\n/Font <<\n/F1 <<\n/Type /Font\n"
            b"/Subtype /Type1\n/BaseFont /Helvetica\n>>\n>>\n>>\n>>\nendobj\n"
            b"4 0 obj\n<<\n/Length 44\n>>\nstream\nBT\n/F1 12 Tf\n100 700 Td\n"
            b"(PDF Data) Tj\nET\nendstream\nendobj\nxref\n0 5\n"
            b"0000000000 65535 f\n0000000009 00000 n\n0000000058 00000 n\n"
            b"0000000115 00000 n\n0000000267 00000 n\ntrailer\n<<\n"
            b"/Size 5\n/Root 1 0 R\n>>\nstartxref\n359\n%%EOF\n"
        )
        return pdf_file

    @pytest.fixture
    def sample_docx_file(self, tmp_path: Path) -> Path:
        from docx import Document

        docx_file = tmp_path / "sample.docx"
        doc = Document()
        doc.add_paragraph("DOCX Data")
        doc.add_paragraph("Sample content from docx")
        doc.save(docx_file)
        return docx_file

    def test_initialization(self, loader: DataLoader) -> None:
        assert loader is not None
        assert len(loader._extractors) == 3

    def test_supported_extensions(self, loader: DataLoader) -> None:
        extensions = loader.supported_extensions
        assert ".pdf" in extensions
        assert ".txt" in extensions
        assert ".md" in extensions
        assert ".markdown" in extensions
        assert ".docx" in extensions

    def test_can_handle(self, loader: DataLoader, tmp_path: Path) -> None:
        txt_file = tmp_path / "test.txt"
        txt_file.touch()
        assert loader.can_handle(txt_file)

        pdf_file = tmp_path / "test.pdf"
        pdf_file.touch()
        assert loader.can_handle(pdf_file)

        docx_file = tmp_path / "test.docx"
        docx_file.touch()
        assert loader.can_handle(docx_file)

        unsupported = tmp_path / "test.xyz"
        unsupported.touch()
        assert not loader.can_handle(unsupported)

    def test_load_single_text_file(self, loader: DataLoader, sample_text_file: Path) -> None:
        result = loader.load(sample_text_file)
        assert isinstance(result, LoadResult)
        assert result.source_count == 1
        assert "Sample text content" in result.unified_text
        assert result.total_chars > 0
        assert result.total_words > 0
        assert not result.has_failures

    def test_load_single_markdown_file(self, loader: DataLoader, sample_md_file: Path) -> None:
        result = loader.load(sample_md_file)
        assert result.source_count == 1
        assert "Markdown" in result.unified_text
        assert "Content here" in result.unified_text

    def test_load_single_pdf_file(self, loader: DataLoader, sample_pdf_file: Path) -> None:
        result = loader.load(sample_pdf_file)
        assert result.source_count == 1
        assert "PDF Data" in result.unified_text

    def test_load_single_docx_file(self, loader: DataLoader, sample_docx_file: Path) -> None:
        result = loader.load(sample_docx_file)
        assert result.source_count == 1
        assert "DOCX Data" in result.unified_text

    def test_load_multiple_files(
        self,
        loader: DataLoader,
        sample_text_file: Path,
        sample_md_file: Path,
    ) -> None:
        result = loader.load([sample_text_file, sample_md_file])
        assert result.source_count == 2
        assert "Sample text content" in result.unified_text
        assert "Markdown" in result.unified_text

    def test_load_with_custom_separator(
        self,
        loader: DataLoader,
        sample_text_file: Path,
        sample_md_file: Path,
    ) -> None:
        result = loader.load([sample_text_file, sample_md_file], separator="\n\n===\n\n")
        assert "\n\n===\n\n" in result.unified_text

    def test_load_raw_text_string(self, loader: DataLoader) -> None:
        raw_text = "This is raw text input"
        result = loader.load(raw_text)
        assert result.source_count == 1
        assert raw_text in result.unified_text

    def test_load_mixed_strings_and_paths(self, loader: DataLoader, sample_text_file: Path) -> None:
        result = loader.load(["Raw text here", sample_text_file])
        assert result.source_count == 2
        assert "Raw text here" in result.unified_text
        assert "Sample text content" in result.unified_text

    def test_load_from_directory(
        self,
        loader: DataLoader,
        sample_text_file: Path,
        sample_md_file: Path,
    ) -> None:
        result = loader.load(sample_text_file.parent, recursive=True)
        assert result.source_count >= 2

    def test_load_from_directory_non_recursive(self, loader: DataLoader, tmp_path: Path) -> None:
        subdir = tmp_path / "subdir"
        subdir.mkdir()
        (subdir / "nested.txt").write_text("nested content")
        (tmp_path / "root.txt").write_text("root content")

        result = loader.load(tmp_path, recursive=False)
        assert "root content" in result.unified_text
        assert "nested content" not in result.unified_text

    def test_load_empty_directory(self, loader: DataLoader, tmp_path: Path) -> None:
        empty_dir = tmp_path / "empty"
        empty_dir.mkdir()
        result = loader.load(empty_dir)
        assert result.source_count == 0
        assert result.has_failures

    def test_load_file_not_found_skip_failures(self, loader: DataLoader, tmp_path: Path) -> None:
        nonexistent = tmp_path / "nonexistent.txt"
        result = loader.load(nonexistent, skip_failures=True)
        assert result.source_count == 0
        assert result.has_failures
        assert len(result.failed_sources) == 1

    def test_load_file_not_found_raise(self, loader: DataLoader, tmp_path: Path) -> None:
        nonexistent = tmp_path / "nonexistent.txt"
        with pytest.raises(FileNotFoundError):
            loader.load(nonexistent, skip_failures=False)

    def test_load_unsupported_file_skip_failures(self, loader: DataLoader, tmp_path: Path) -> None:
        unsupported = tmp_path / "test.xyz"
        unsupported.write_text("content")
        result = loader.load(unsupported, skip_failures=True)
        assert result.source_count == 0

    def test_load_invalid_file_skip_failures(self, loader: DataLoader, tmp_path: Path) -> None:
        invalid_pdf = tmp_path / "invalid.pdf"
        invalid_pdf.write_text("not a pdf")
        result = loader.load(invalid_pdf, skip_failures=True)
        assert result.has_failures

    def test_load_invalid_file_raise(self, loader: DataLoader, tmp_path: Path) -> None:
        invalid_pdf = tmp_path / "invalid.pdf"
        invalid_pdf.write_text("not a pdf")
        with pytest.raises(ExtractionError):
            loader.load(invalid_pdf, skip_failures=False)

    def test_load_single_method(self, loader: DataLoader, sample_text_file: Path) -> None:
        text = loader.load_single(sample_text_file)
        assert isinstance(text, str)
        assert "Sample text content" in text

    def test_load_single_with_string_path(self, loader: DataLoader, sample_text_file: Path) -> None:
        text = loader.load_single(str(sample_text_file))
        assert "Sample text content" in text

    def test_load_single_file_not_found(self, loader: DataLoader, tmp_path: Path) -> None:
        with pytest.raises(FileNotFoundError):
            loader.load_single(tmp_path / "nonexistent.txt")

    def test_load_result_properties(self, loader: DataLoader, tmp_path: Path) -> None:
        txt_file = tmp_path / "test.txt"
        txt_file.write_text("Hello world test")

        result = loader.load(txt_file)
        assert result.total_chars == len("Hello world test")
        assert result.total_words == 3
        assert result.source_count == 1
        assert not result.has_failures

    def test_load_skips_empty_files(self, loader: DataLoader, tmp_path: Path) -> None:
        empty_file = tmp_path / "empty.txt"
        empty_file.write_text("")
        non_empty_file = tmp_path / "non_empty.txt"
        non_empty_file.write_text("content")

        result = loader.load([empty_file, non_empty_file])
        assert result.source_count == 1

    def test_collect_files_single_file(self, loader: DataLoader, sample_text_file: Path) -> None:
        files = loader._collect_files(sample_text_file)
        assert len(files) == 1
        assert files[0] == sample_text_file

    def test_collect_files_unsupported_file(self, loader: DataLoader, tmp_path: Path) -> None:
        unsupported = tmp_path / "test.xyz"
        unsupported.touch()
        files = loader._collect_files(unsupported)
        assert len(files) == 0

    def test_get_extractor_returns_correct_extractor(
        self, loader: DataLoader, tmp_path: Path
    ) -> None:
        txt_file = tmp_path / "test.txt"
        txt_file.touch()
        extractor = loader._get_extractor(txt_file)
        assert extractor is not None
        assert isinstance(extractor, TextExtractor)

        pdf_file = tmp_path / "test.pdf"
        pdf_file.touch()
        extractor = loader._get_extractor(pdf_file)
        assert extractor is not None
        assert isinstance(extractor, PDFExtractor)

        docx_file = tmp_path / "test.docx"
        docx_file.touch()
        extractor = loader._get_extractor(docx_file)
        assert extractor is not None
        assert isinstance(extractor, DocxExtractor)

    def test_get_extractor_returns_none_for_unsupported(
        self, loader: DataLoader, tmp_path: Path
    ) -> None:
        unsupported = tmp_path / "test.xyz"
        unsupported.touch()
        extractor = loader._get_extractor(unsupported)
        assert extractor is None


class TestLoadResult:
    """Tests for LoadResult dataclass."""

    def test_create_result(self) -> None:
        result = LoadResult(unified_text="Combined text")
        assert result.unified_text == "Combined text"
        assert result.source_count == 0
        assert not result.has_failures

    def test_result_with_sources(self, tmp_path: Path) -> None:
        source1 = ExtractionResult(text="Text 1", source_path=tmp_path / "1.txt")
        source2 = ExtractionResult(text="Text 2", source_path=tmp_path / "2.txt")
        result = LoadResult(unified_text="Text 1\n\nText 2", sources=[source1, source2])
        assert result.source_count == 2
        assert not result.has_failures

    def test_result_with_failures(self, tmp_path: Path) -> None:
        result = LoadResult(
            unified_text="",
            failed_sources=[(tmp_path / "fail.txt", "Failed to read")],
        )
        assert result.has_failures
        assert len(result.failed_sources) == 1

    def test_total_chars_and_words(self) -> None:
        text = "Hello world this is a test"
        result = LoadResult(unified_text=text)
        assert result.total_chars == len(text)
        assert result.total_words == 6
